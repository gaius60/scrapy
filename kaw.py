import docx
import xlsxwriter

# Load the word document
doc = docx.Document('words.docx')

# Initialize an empty dictionary to store word frequencies
word_freq = {}

# Loop through each paragraph in the document
for para in doc.paragraphs:
    # Split the text in the paragraph by commas to extract individual words
    words = para.text.split(',')
    
    # Loop through each word
    for word in words:
        # Strip any leading or trailing whitespace from the word
        word = word.strip()
        
        # If the word is already in the dictionary, increment its frequency count
        if word in word_freq:
            word_freq[word] += 1
        # Otherwise, add the word to the dictionary with a frequency of 1
        else:
            word_freq[word] = 1

# Sort the dictionary by frequency in descending order and take the top 100 words
top_words = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)[:100]

# Create a new Excel workbook and worksheet
workbook = xlsxwriter.Workbook('top_words.xlsx')
worksheet = workbook.add_worksheet()

# Write the top 100 words and frequencies to the worksheet
row = 0
for i, (word, freq) in enumerate(top_words):
    worksheet.write(row, 0, i+1)  # Write the rank
    worksheet.write(row, 1, word)  # Write the word
    worksheet.write(row, 2, freq)  # Write the frequency
    row += 1

# Close the workbook
workbook.close()

print('Top 100 used words and their frequencies have been written to top_words.xlsx')

